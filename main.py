from pulp import *
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

# =====================================
# PODACI
# =====================================

agents = [f"A{i}" for i in range(1,11)]
supervisors = [f"S{i}" for i in range(1,5)]

employees = agents + supervisors

days = ["Ponedjeljak","Utorak","Srijeda","Četvrtak","Petak","Subota","Nedjelja"]

weekday_shifts = ["SUP_6_14","AG_7_15","AG_8_16","AG_10_18","AG_11_19","SUP_14_22"]
weekend_shifts = ["WK_6_14","WK_14_22"]
all_shifts = weekday_shifts + weekend_shifts

# =====================================
# MODEL
# =====================================

model = LpProblem("CallCenter", LpMinimize)

x = LpVariable.dicts("work",(employees,days,all_shifts),cat="Binary")

# minimalni broj smjena
model += lpSum(x[e][d][s] for e in employees for d in days for s in all_shifts)

# jedna smjena dnevno
for e in employees:
    for d in days:
        model += lpSum(x[e][d][s] for s in all_shifts) <= 1

# max 5 smjena tjedno
for e in employees:
    model += lpSum(x[e][d][s] for d in days for s in all_shifts) <= 5

# radni dani
for d in days[:5]:

    model += lpSum(x[sv][d]["SUP_6_14"] for sv in supervisors) == 1
    model += lpSum(x[sv][d]["SUP_14_22"] for sv in supervisors) == 1

    model += lpSum(x[e][d]["AG_7_15"] for e in employees) == 1
    model += lpSum(x[e][d]["AG_10_18"] for e in employees) == 1
    model += lpSum(x[e][d]["AG_11_19"] for e in employees) == 1
    model += lpSum(x[e][d]["AG_8_16"] for e in employees) == 7

# vikend
for d in days[5:]:
    model += lpSum(x[e][d]["WK_6_14"] for e in employees) == 1
    model += lpSum(x[e][d]["WK_14_22"] for e in employees) == 1

# agent ne može biti supervizor
for a in agents:
    for d in days:
        model += x[a][d]["SUP_6_14"] == 0
        model += x[a][d]["SUP_14_22"] == 0

# zabrani radne smjene vikendom i obrnuto
for d in days[:5]:
    model += lpSum(x[e][d]["WK_6_14"] + x[e][d]["WK_14_22"] for e in employees) == 0

for d in days[5:]:
    model += lpSum(x[e][d][s] for e in employees for s in weekday_shifts) == 0

model.solve(PULP_CBC_CMD(msg=True))

print("STATUS:", LpStatus[model.status])

# =====================================
# EXCEL
# =====================================

rows = []

for d in days:
    row = {"Dan": d}

    for s in all_shifts:
        assigned = []
        for e in employees:
            if value(x[e][d][s]) == 1:
                assigned.append(e)

        row[s] = ", ".join(assigned)

    rows.append(row)

df = pd.DataFrame(rows)

file_name = "Raspored_Call_Centar.xlsx"
with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Raspored", index=False)

wb = load_workbook(file_name)
ws = wb["Raspored"]

for cell in ws[1]:
    cell.font = Font(bold=True)

wb.save(file_name)

print("Generiran:", file_name)
